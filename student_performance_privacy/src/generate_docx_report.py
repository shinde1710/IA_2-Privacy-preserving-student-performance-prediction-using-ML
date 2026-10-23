"""
generate_docx_report.py
-----------------------
Generates a formatted Microsoft Word (.docx) document for the IA 1 technical report.
Includes:
- Academic document styling
- Clean tables with borders and styled headers
- Embedded figures from the experimental results
- Complete 17-section structure
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def create_docx_report(base_dir: str):
    doc = Document()
    
    # Page setup: Standard margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Privacy-Preserving Student Performance Prediction\nUsing Machine Learning")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    # Subtitle / Metadata
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Internal Assessment 1 (IA 1) Preliminary Technical Report\n"
                            "Subject: Security and Privacy for Machine Learning | Semester VII B.Tech\n"
                            "Student: Aryan | Academic Year: 2026-2027")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Helper functions
    def add_sec_heading(num, text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(f"{num}. {text}")
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        return h

    def add_sub_heading(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
        return h

    def add_body_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        return p

    # 1. Project Details
    add_sec_heading(1, "Project Details and Title")
    add_body_p("Project Title: Privacy-Preserving Student Performance Prediction Using Machine Learning\n"
               "Project Category: Privacy-Preserving Machine Learning, Educational Data Mining, Differential Privacy, Responsible AI.\n"
               "Primary Learning Algorithm: Pruned Decision Tree Classifier (with Logistic Regression comparison).\n"
               "Privacy Paradigms: Data Anonymization (k-anonymity, Quasi-Identifier Uniqueness) and Local Differential Privacy (Client-side Bounded Laplace Mechanism).\n"
               "Target Feature: Final mathematics outcome (G3 >= 10 -> Pass, G3 < 10 -> Fail).\n"
               "Data Leakage Safeguard: Strict exclusion of intermediate exam grades G1 and G2.")

    # 2. Abstract
    add_sec_heading(2, "Abstract / Executive Summary")
    add_body_p("Educational data mining models offer substantial potential for early scholastic interventions by predicting which students are at risk of academic failure. However, educational records encapsulate highly sensitive personal information, including family cohabitation status, parental education levels, home address types, study habits, and disciplinary absences. Training conventional machine learning models directly on raw, centralized student records introduces severe privacy risks, including linkage-based re-identification and sensitive attribute inference.\n\n"
               "This project investigates the empirical trade-off between privacy guarantees and downstream classification utility for student performance prediction on the benchmark UCI Student Performance dataset (N=395). To reflect realistic institutional deployment, we eliminate target leakage by excluding intermediate examination grades (G1, G2), compelling the model to predict outcomes using demographic, behavioral, and environmental indicators.\n\n"
               "We evaluate three modeling regimes: (1) a baseline interpretable Decision Tree achieving 65.82% accuracy and 78.05% F1-score, while exposing 27.59% of students as uniquely re-identifiable; (2) a generalized, de-identified feature space reducing uniqueness to 1.27% while improving accuracy to 68.35% and maintaining a 77.88% F1-score; and (3) a client-side Local Differential Privacy framework applying bounded Laplace perturbation to study time across privacy budgets epsilon in [0.1, 10.0]. The findings demonstrate that moderate anonymization and localized feature perturbation maintain strong model utility while substantially curbing individual exposure risks.")

    # 3. Problem Statement
    add_sec_heading(3, "Problem Statement")
    add_body_p("Modern educational institutions collect multifaceted student data to identify at-risk students. However, transmitting and storing unencrypted, granular student records creates grave vulnerabilities:\n"
               "1. Linkage Attacks: External adversaries possessing auxiliary demographic datasets (e.g., voter registration lists) can join records to re-identify individual students.\n"
               "2. Internal Curiosity and Untrusted Aggregators: Central administrators or third-party cloud data brokers can observe intimate details including parental cohabitation, alcohol consumption, and study habits.\n"
               "3. Model Inversion and Membership Inference: Machine learning models trained on unprotected records can memorize specific student records, exposing private details through prediction logits.\n\n"
               "The central research question is: How can an educational institution train an effective, explainable machine learning model to predict student failure early in the academic cycle while providing verifiable privacy safeguards for sensitive student attributes?")

    # 4. Motivation and Significance
    add_sec_heading(4, "Motivation and Significance")
    add_body_p("1. Regulatory Compliance: Regulatory frameworks including FERPA (US) and GDPR (EU) legally mandate strict privacy safeguards, data minimization, and protection against unauthorized re-identification.\n"
               "2. Student Trust and Ethical AI: If students fear surveillance or disclosure of personal lifestyle habits, they may report inaccurate data or refuse participation, undermining institutional analytics.\n"
               "3. Early Intervention without Data Leakage: Naive predictive models often include earlier test grades (G1, G2), achieving misleadingly high accuracy (>90%). However, waiting for mid-semester grades prevents early academic support. A legitimate early-intervention model must rely on pre-examination demographic and behavioral attributes.")

    # 5. Objectives and Expected Outcomes
    add_sec_heading(5, "Objectives and Expected Outcomes")
    add_body_p("Objectives:\n"
               "1. Construct a leakage-free prediction pipeline predicting mathematics outcome (Pass/Fail) without prior exam scores.\n"
               "2. Train a pruned, easily explainable Decision Tree classifier as the primary academic baseline.\n"
               "3. Implement Experiment B: Quantify re-identification vulnerability on raw quasi-identifiers and evaluate generalization/suppression transformations.\n"
               "4. Implement Experiment C: Apply a mathematically sound client-side Local Differential Privacy (LDP) Laplace mechanism on study time.\n"
               "5. Implement Experiment D: Formulate the privacy-utility trade-off curve across privacy budgets epsilon in [0.1, 10.0].\n\n"
               "Expected Outcomes: A reproducible, modular Python project and Jupyter Notebook demonstrating that de-identification reduces re-identification risk by 95% while preserving predictive utility.")

    # 6. Scope and Limitations
    add_sec_heading(6, "Scope and Limitations")
    add_body_p("Scope: UCI Student Performance Mathematics dataset (395 records, 33 attributes); binary classification (G3 >= 10); Decision Tree primary model; Logistic Regression baseline; k-anonymity analysis; single-attribute Local DP.\n\n"
               "Limitations: Multi-party federated learning and DP-SGD across all features are deferred to IA 2. Findings represent a Portuguese secondary school sample and should not be generalized across disparate cultural contexts without recalibration.")

    # 7. Conceptual Background
    add_sec_heading(7, "Conceptual Background")
    add_body_p("Decision Trees: Partition the input feature space recursively by minimizing Gini impurity. In educational settings, Decision Trees offer superior transparency because decision paths can be directly read as policy rules (e.g., if past failures <= 0 and absences <= 15, then Pass).\n\n"
               "k-Anonymity: Formulated by Sweeney (2002), k-anonymity requires that each combination of quasi-identifiers occurs at least k times in the released dataset. An equivalence class with k=1 indicates a uniquely identifiable student.\n\n"
               "Differential Privacy: Central DP guarantees that the presence or absence of any single student has a bounded impact on the model output. Local DP perturbs data on the student device prior to collection, protecting individuals even against an untrusted central aggregator. The Post-Processing Theorem guarantees that clipping noisy outputs does not compromise the epsilon-LDP guarantee.")

    # 8. Mapping to All Five Syllabus Modules
    add_sec_heading(8, "Mapping to All Five Syllabus Modules")
    
    t_mod = doc.add_table(rows=6, cols=5)
    t_mod.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Syllabus Module", "Course Concept", "Project Implementation", "Empirical Evidence", "Limitations"]
    
    for i, h in enumerate(headers):
        cell = t_mod.cell(0, i)
        cell.text = h
        set_cell_background(cell, "1B365D")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    mod_data = [
        ("Module 1: Privacy Considerations in ML", "Sensitive academic data, inference risks, data leakage.", "Strict exclusion of G1 and G2; demographic quasi-identifier audit.", "Feature matrix reduced from 33 to 30; class distribution plot.", "Requires institutional review board and consent."),
        ("Module 2: Differential Privacy for ML", "Central DP, privacy budget epsilon, noise calibration.", "Derivation of Laplace scale lambda = Delta_f / epsilon.", "Sensitivity Delta_f = 3; budget analysis across eps in [0.1, 10.0].", "Central DP assumes trusted curator; DP-SGD in IA 2."),
        ("Module 3: Local Differential Privacy for ML", "Untrusted aggregator, client perturbation, feature distortion.", "Client-side bounded Laplace perturbation on studytime with clipping.", "MAE distortion curve: 1.4770 (eps=0.1) to 0.2403 (eps=10.0).", "High noise at small epsilon introduces individual variance."),
        ("Module 4: Privacy-Preserving Data Mining", "Quasi-identifiers, k-anonymity, generalization, suppression.", "Equivalence class analysis; binning age, Medu, Fedu; suppressing school.", "Unique students reduced from 109 (27.6%) to 5 (1.3%); 68.35% accuracy.", "Vulnerable to homogeneity attacks; lacks DP bound."),
        ("Module 5: Compressive Privacy for ML", "Representation learning, information bottleneck, sensitive suppression.", "Coarse binning and tree pruning suppressing demographic specifics.", "Downstream model maintains 77.88% F1-score on compressed features.", "PCA alone preserves variance, not privacy bounds.")
    ]
    
    for r_idx, row in enumerate(mod_data, start=1):
        bg = "F4F6F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = t_mod.cell(r_idx, c_idx)
            cell.text = val
            set_cell_background(cell, bg)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 9. Dataset Provenance
    add_sec_heading(9, "Dataset Provenance and Responsible Data Use")
    add_body_p("Dataset Source: UCI Machine Learning Repository (Cortez & Silva, 2008). CC BY 4.0 License.\n"
               "Cohort: 395 students in Portuguese secondary school mathematics classes.\n"
               "Data Integrity: Verified 0 missing entries, 0 duplicate records.\n"
               "Ethical Protocol: Used exclusively for privacy-preserving ML research. All direct identifiers are absent.")

    # 10. Methodology
    add_sec_heading(10, "Methodology and Workflow")
    add_body_p("The pipeline encompasses data ingestion, leakage-free binarization, stratified 80/20 train/test splitting, baseline Decision Tree training, k-anonymity auditing/generalization, client-side Local DP bounded Laplace perturbation, and multi-metric evaluation.")

    # 11. Initial Implementation Details
    add_sec_heading(11, "Initial Implementation Details")
    add_body_p("Implementation utilizes Python 3.12, scikit-learn, pandas, numpy, and matplotlib. Modular architecture comprises: data_preprocessing.py, baseline_model.py, anonymization.py, local_differential_privacy.py, and evaluation.py.")

    # 12. Experimental Setup
    add_sec_heading(12, "Experimental Setup")
    add_body_p("Reproducibility: Fixed random seed 42 across all splits, noise generations, and estimators.\n"
               "Decision Tree: max_depth=3, min_samples_leaf=5, Gini criterion.\n"
               "Logistic Regression: max_iter=1000, L-BFGS solver.\n"
               "Local DP: Laplace mechanism on studytime (domain [1, 4], Delta_f = 3) across epsilons [0.1, 0.5, 1.0, 2.0, 5.0, 10.0].")

    # 13. Evaluation Metrics
    add_sec_heading(13, "Evaluation Metrics")
    add_body_p("Classification performance: Accuracy, Precision, Recall, F1-Score, and full Confusion Matrix (TN, FP, FN, TP). Distortion metrics: Mean Absolute Error (MAE) and Root Mean Squared Error (RMSE).")

    # 14. Actual Preliminary Results
    add_sec_heading(14, "Actual Preliminary Results")
    
    # Class distribution figure
    class_fig = os.path.join(base_dir, "figures", "class_distribution.png")
    if os.path.exists(class_fig):
        doc.add_picture(class_fig, width=Inches(5.0))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 1: Target Class Distribution (Pass: 265, Fail: 130)")
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        
    add_body_p("Baseline Model Outputs:\n"
               "- Decision Tree: Accuracy 65.82%, Precision 68.57%, Recall 90.57%, F1-score 78.05%, Confusion Matrix [[4, 22], [5, 48]].\n"
               "- Logistic Regression: Accuracy 65.82%, Precision 71.67%, Recall 81.13%, F1-score 76.11%, Confusion Matrix [[9, 17], [10, 43]].")

    # Confusion Matrix figure
    cm_fig = os.path.join(base_dir, "figures", "confusion_matrix.png")
    if os.path.exists(cm_fig):
        doc.add_picture(cm_fig, width=Inches(4.5))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 2: Baseline Decision Tree Confusion Matrix")
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True

    add_sub_heading("Experiment B: Data Anonymization & k-Anonymity Audit")
    add_body_p("Raw quasi-identifiers (school, sex, age, address, Medu, Fedu) exposed 109 students (27.59%) as completely unique (k=1). After generalization (suppressing school, binning age, Medu, Fedu, absences), uniquely identifiable students dropped to 5 (1.27%)—a 95.4% reduction in re-identification vulnerability. Downstream Decision Tree accuracy improved to 68.35% with an F1-score of 77.88%.")

    add_sub_heading("Experiment C: Local Differential Privacy Distortion")
    add_body_p("Client-side Laplace perturbation on studytime produced clear empirical trade-offs between epsilon and distortion error:")

    # LDP Table
    t_ldp = doc.add_table(rows=7, cols=5)
    t_ldp.alignment = WD_TABLE_ALIGNMENT.CENTER
    ldp_headers = ["Epsilon (eps)", "Privacy Level", "MAE Distortion", "RMSE Distortion", "Downstream F1"]
    for i, h in enumerate(ldp_headers):
        cell = t_ldp.cell(0, i)
        cell.text = h
        set_cell_background(cell, "1B365D")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    ldp_rows = [
        ("0.1", "Strong", "1.4770", "1.7589", "0.7805"),
        ("0.5", "Strong", "1.2905", "1.5750", "0.7805"),
        ("1.0", "Moderate", "1.1030", "1.3880", "0.7805"),
        ("2.0", "Moderate", "0.8305", "1.0890", "0.7805"),
        ("5.0", "Weak", "0.4520", "0.6295", "0.7805"),
        ("10.0", "Weak", "0.2403", "0.3445", "0.7805")
    ]
    for r_idx, row in enumerate(ldp_rows, start=1):
        bg = "F4F6F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = t_ldp.cell(r_idx, c_idx)
            cell.text = val
            set_cell_background(cell, bg)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Figures: Distortion and Trade-off
    dist_fig = os.path.join(base_dir, "figures", "feature_distortion.png")
    if os.path.exists(dist_fig):
        doc.add_picture(dist_fig, width=Inches(5.0))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 3: Feature Distortion Error vs. Privacy Budget (Epsilon)")
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True

    add_sub_heading("Experiment D: Consolidated Privacy-Utility Comparison")
    
    # Final Comparison Table
    t_comp = doc.add_table(rows=7, cols=6)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_headers = ["Method", "Privacy Mechanism", "Accuracy", "Precision", "Recall", "F1-Score"]
    for i, h in enumerate(c_headers):
        cell = t_comp.cell(0, i)
        cell.text = h
        set_cell_background(cell, "1B365D")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    comp_rows = [
        ("Baseline Decision Tree", "None (Raw features)", "0.6582", "0.6857", "0.9057", "0.7805"),
        ("Baseline Logistic Reg", "None (Linear model)", "0.6582", "0.7167", "0.8113", "0.7611"),
        ("Anonymized Decision Tree", "Generalization + suppression", "0.6835", "0.7333", "0.8302", "0.7788"),
        ("LDP Decision Tree (eps=0.5)", "Laplace on studytime", "0.6582", "0.6857", "0.9057", "0.7805"),
        ("LDP Decision Tree (eps=1.0)", "Laplace on studytime", "0.6582", "0.6857", "0.9057", "0.7805"),
        ("LDP Decision Tree (eps=5.0)", "Laplace on studytime", "0.6582", "0.6857", "0.9057", "0.7805")
    ]
    for r_idx, row in enumerate(comp_rows, start=1):
        bg = "F4F6F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = t_comp.cell(r_idx, c_idx)
            cell.text = val
            set_cell_background(cell, bg)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    model_fig = os.path.join(base_dir, "figures", "model_comparison.png")
    if os.path.exists(model_fig):
        doc.add_picture(model_fig, width=Inches(5.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 4: Performance Comparison across Baseline and Privacy Models")
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True

    # 15. Observations
    add_sec_heading(15, "Observations and Limitations")
    add_body_p("1. Model Resilience: The Decision Tree relies on past failures and absences as top split features, making it highly robust against localized noise on study time.\n"
               "2. Anonymization Limitations: While k-anonymity reduces uniqueness, it does not defend against homogeneity attacks. Formal guarantees require Differential Privacy.\n"
               "3. Realistic Early Prediction: Excluding G1 and G2 ensures the model can be deployed before exams, providing authentic academic utility.")

    # 16. Work Remaining
    add_sec_heading(16, "Work Remaining for IA 2")
    add_body_p("For IA 2, future work will implement: (1) Differentially Private SGD (DP-SGD) and DP Decision Forests; (2) Compressive privacy representation learning via adversarial networks; (3) Multi-attribute Local DP for both categorical and numerical features.")

    # 17. References
    add_sec_heading(17, "References")
    add_body_p("1. Cortez, P., & Silva, A. (2008). Using data mining to predict secondary school student performance. FUBUTEC 2008.\n"
               "2. Dwork, C., & Roth, A. (2014). The algorithmic foundations of differential privacy. Foundations and Trends in TCS.\n"
               "3. Sweeney, L. (2002). k-anonymity: A model for protecting privacy. Int. J. Uncertain. Fuzziness Knowl.-Based Syst.\n"
               "4. Machanavajjhala, A. et al. (2007). l-diversity: Privacy beyond k-anonymity. ACM TKDD.\n"
               "5. Duchi, J. C. et al. (2013). Local privacy and statistical minimax rates. IEEE FOCS.")

    out_docx = os.path.join(base_dir, "reports", "IA1_Preliminary_Report.docx")
    doc.save(out_docx)
    print(f"[OK] Generated Word report: {out_docx}")


if __name__ == "__main__":
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    create_docx_report(base_dir)
